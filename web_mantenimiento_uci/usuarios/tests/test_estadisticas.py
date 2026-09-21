"""Dashboard, conteos por mes y exportación a Excel."""
import io
from datetime import datetime

import openpyxl
from django.utils import timezone

from usuarios import estadisticas
from usuarios.models import Incidencia

from .base import SGUMTestCase


def en(anio, mes, dia=15):
    return timezone.make_aware(datetime(anio, mes, dia, 12, 0))


class ConteoPorMesTest(SGUMTestCase):

    def setUp(self):
        self.anio = timezone.localdate().year
        self.crear_incidencia(fecha=en(self.anio, 1))
        self.crear_incidencia(fecha=en(self.anio, 3, 2))
        self.crear_incidencia(fecha=en(self.anio, 3, 20), estado='resuelto')
        self.crear_incidencia(fecha=en(self.anio, 3, 28), estado='en_proceso')
        self.crear_incidencia(fecha=en(self.anio - 1, 3))  # otro año: no cuenta
        self.esperado = [1, 0, 3, 0, 0, 0, 0, 0, 0, 0, 0, 0]

    def test_funcion_compartida_agrupa_por_mes(self):
        self.assertEqual(estadisticas.incidencias_por_mes(self.anio), self.esperado)

    def test_dashboard_y_exportacion_dan_lo_mismo(self):
        self.entrar(self.admin)
        contexto = self.client.get(self.url('reportes')).context
        self.assertEqual(contexto['incidencias_data'], self.esperado)

        respuesta = self.client.get(self.url('exportar_dashboard'))
        libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
        hoja = libro['Por Mes']
        valores = [hoja.cell(row=fila, column=2).value for fila in range(2, 14)]
        self.assertEqual(valores, self.esperado)

    def test_conteo_por_estado(self):
        conteos = estadisticas.contar_por_estado()
        self.assertEqual((conteos['pendiente'], conteos['en_proceso'], conteos['resuelto'],
                          conteos['total']), (3, 1, 1, 5))
        self.assertEqual(conteos['abiertas'], 4)
        self.assertEqual(conteos['sin_tecnico'], 4)

    def test_conteo_por_tipo_incluye_los_diez_tipos(self):
        filas = estadisticas.contar_por_tipo()
        self.assertEqual(len(filas), 10)
        self.assertEqual(next(f for f in filas if f['codigo'] == 'plomeria')['cantidad'], 5)


class DashboardTest(SGUMTestCase):

    def test_filtro_por_mes(self):
        anio = timezone.localdate().year
        self.crear_incidencia(fecha=en(anio, 2), estado='resuelto')
        self.crear_incidencia(fecha=en(anio, 2, 20))
        self.crear_incidencia(fecha=en(anio, 5))
        self.entrar(self.almacenero)
        contexto = self.client.get(self.url('reportes'), {'mesAnio': f'{anio}-02'}).context
        self.assertEqual(contexto['totalReportes'], 2)
        self.assertEqual(contexto['reporte_resuelto'], 1)
        self.assertEqual(contexto['reporte_pendiente'], 1)
        self.assertEqual(len(list(contexto['tableReporte'])), 2)
        self.assertEqual(contexto['mesAnio'], f'{anio}-02')

    def test_mes_invalido_no_rompe(self):
        self.crear_incidencia()
        self.entrar(self.admin)
        for valor in ('basura', '2025-13', '2025-00', '99999-01', '2025'):
            with self.subTest(valor=valor):
                respuesta = self.client.get(self.url('reportes'), {'mesAnio': valor})
                self.assertEqual(respuesta.status_code, 200)
                self.assertEqual(respuesta.context['totalReportes'], 1)
                self.assertEqual(self.mensajes(respuesta)[0][0], 'warning')

    def test_expone_estado_y_tipo_para_los_graficos(self):
        self.crear_incidencia(tipo='gas')
        self.entrar(self.admin)
        contexto = self.client.get(self.url('reportes')).context
        self.assertEqual([f['codigo'] for f in contexto['por_estado']],
                         ['pendiente', 'en_proceso', 'resuelto'])
        self.assertEqual(len(contexto['por_tipo']), 10)
        self.assertEqual(len(contexto['meses']), 12)


class ExportacionTest(SGUMTestCase):

    def exportar(self):
        self.entrar(self.admin)
        respuesta = self.client.get(self.url('exportar_dashboard'))
        self.assertEqual(respuesta.status_code, 200)
        return respuesta, openpyxl.load_workbook(io.BytesIO(respuesta.content))

    def test_es_un_xlsx_con_las_hojas_esperadas(self):
        self.crear_incidencia(estado='en_proceso', tecnico_asignado=self.personal)
        respuesta, libro = self.exportar()
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        self.assertIn('attachment', respuesta['Content-Disposition'])
        self.assertEqual(libro.sheetnames, ['Estados', 'Por Mes', 'Por Tipo', 'Listado'])
        estados = {libro['Estados'].cell(row=f, column=1).value: libro['Estados'].cell(row=f, column=2).value
                   for f in range(2, 5)}
        self.assertEqual(estados, {'Pendiente': 0, 'Resuelto': 0, 'En proceso': 1})

    def test_listado_usa_etiquetas_y_datos_completos(self):
        self.crear_incidencia(estado='en_proceso', tecnico_asignado=self.personal,
                              tipo='mantenimiento_equipos', prioridad='3')
        _, libro = self.exportar()
        fila = [c.value for c in libro['Listado'][2]]
        self.assertEqual(fila[2], 'Mantenimiento de equipos')
        self.assertEqual(fila[4], 'En proceso')
        self.assertEqual(fila[5], 'Alta')
        self.assertEqual(fila[6], 'No')
        self.assertEqual(fila[8], self.solicitante.username)
        self.assertEqual(fila[9], self.tecnico.username)

    def test_una_descripcion_no_se_ejecuta_como_formula(self):
        self.crear_incidencia(descripcion='=HYPERLINK("http://x.example","clic")')
        _, libro = self.exportar()
        celda = libro['Listado'].cell(row=2, column=4)
        self.assertNotEqual(celda.data_type, 'f')
        self.assertEqual(celda.value, '=HYPERLINK("http://x.example","clic")')

    def test_caracteres_de_control_no_rompen_la_exportacion(self):
        self.crear_incidencia(descripcion='hola\x07mundo')
        self.exportar()

    def test_exportacion_sin_datos(self):
        _, libro = self.exportar()
        self.assertEqual(libro['Listado'].max_row, 1)

    def test_solo_administrador_y_almacenero(self):
        for usuario, esperado in ((self.solicitante, 403), (self.tecnico, 403),
                                  (self.almacenero, 200), (self.admin, 200)):
            self.entrar(usuario)
            self.assertEqual(self.client.get(self.url('exportar_dashboard')).status_code, esperado)
